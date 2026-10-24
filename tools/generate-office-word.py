#!/usr/bin/env python3
"""Synthetic Word fixture. Needs python-docx 1.2.0 and Pillow. No external resources."""
import argparse
import io
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Liberation Serif', Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    for style in doc.styles:
        properties = style.element.find(qn('w:pPr'))
        if properties is not None:
            for border in list(properties.findall(qn('w:pBdr'))):
                properties.remove(border)
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    section.header.paragraphs[0].text = 'CAPT — синтетический тест кириллицы'
    footer = section.footer.paragraphs[0]
    footer.add_run('Страница ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.add_paragraph('Проверка вёрстки Word', 'Title')
    doc.add_paragraph('ПЕРВАЯ СТРАНИЦА. Этот документ проверяет сохранение оформления при печати.')
    p = doc.add_paragraph('Обычный текст. ')
    p.add_run('Полужирный. ').bold = True
    p.add_run('Курсив. ').italic = True
    p.add_run('Подчёркнутый. ').underline = True
    p.add_run('Цветной.').font.color.rgb = RGBColor(0x20, 0x40, 0xA0)
    doc.add_paragraph('Первый пункт списка', 'List Number')
    doc.add_paragraph('Второй пункт списка', 'List Number')
    doc.add_paragraph('Маркированный пункт', 'List Bullet')
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f'Ячейка {i + 1} {j + 1}'
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Объединённые ячейки'
    doc.add_paragraph()
    image = Image.new('RGB', (300, 100), 'white')
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((2, 2, 297, 97), outline='black', width=3)
    drawing.rectangle((15, 15, 90, 85), fill='black')
    drawing.ellipse((120, 15, 190, 85), outline='black', width=3)
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    buffer.seek(0)
    doc.add_picture(buffer, width=Mm(90))
    doc.add_page_break()
    doc.add_paragraph('ВТОРАЯ СТРАНИЦА', 'Heading 1')
    doc.add_paragraph('Эта строка должна начинаться после явного разрыва страницы.')
    doc.add_paragraph('Ёж, съешь ещё этих мягких французских булок. 0123456789 ± × ÷.')
    landscape = doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = Mm(297), Mm(210)
    doc.add_paragraph('ТРЕТЬЯ СТРАНИЦА', 'Heading 1')
    doc.add_paragraph('Альбомный раздел A4. Поворот должен быть виден до печати.')
    doc.save(args.output / 'word.docx')


if __name__ == '__main__':
    main()
